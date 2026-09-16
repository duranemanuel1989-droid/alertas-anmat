#!/usr/bin/env python3
"""
Bot del Boletin de ANMAT (Helena - Productos Medicos) -> Telegram.

Lee el listado de productos medicos en https://helena.anmat.gob.ar/Boletin/
y avisa por Telegram cuando aparecen novedades. Ahora revisa LAS DOS solapas
de la pagina:

  - "Registros"      (tabla gvTramites)
  - "Notificaciones" (tabla gvDeclaraciones)  -> declaraciones juradas PM I-II

Cuando hay novedades, en vez de mandar fotos manda UN archivo Excel (.xlsx)
con UNA hoja que junta las dos solapas, con una columna "Tipo"
(Registro / Notificación) y ordenada por Razon Social (para ver juntas todas
las filas de una misma empresa). Cada fila trae: Tipo, Tramite, Fecha, Razon
Social, Nombre, Marca, Modelo/s (texto completo), PM y Expediente.

Solo avisa lo nuevo; el estado ya avisado se guarda en 'vistos_boletin.json'.

Usa los mismos Secrets que el bot de alertas:
  - TELEGRAM_TOKEN
  - TELEGRAM_CHAT_ID
"""

import os
import sys
import json
from datetime import datetime
from pathlib import Path

import requests
from bs4 import BeautifulSoup

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

try:
    from zoneinfo import ZoneInfo
    TZ = ZoneInfo("America/Argentina/Buenos_Aires")
except Exception:
    TZ = None

URL = "https://helena.anmat.gob.ar/Boletin/"
# Las dos solapas de la pagina ya vienen ambas en el HTML inicial (son pestañas
# tipo Bootstrap: #tabRegistros y #tabDeclaraciones), no hace falta postback.
TABLA_REGISTROS = "ctl00_ContentPlaceHolder1_gvTramites"
TABLA_NOTIFICACIONES = "ctl00_ContentPlaceHolder1_gvDeclaraciones"

STATE_FILE = Path("vistos_boletin.json")

TOKEN = os.environ.get("TELEGRAM_TOKEN", "").strip()
CHAT_ID = os.environ.get("TELEGRAM_CHAT_ID", "").strip()

# Columnas (mismo orden en las dos tablas para los indices 0..7):
#   0 Tramite | 1 Fecha Fin | 2 Razon Social | 3 Nombre | 4 Marca
#   5 Modelo/s | 6 PM | 7 Expediente | (8/9 links de documento, no se usan)
COL = {
    "tramite": 0,
    "fecha": 1,
    "razon": 2,
    "nombre": 3,
    "marca": 4,
    "modelo": 5,
    "pm": 6,
    "expediente": 7,
}


def _es_expediente(exp):
    """True si 'exp' parece un expediente ANMAT real (ej. 1-0047-3110-004096-26-3).

    Sirve para descartar las filas basura que la grilla ASP.NET agrega al final
    (el paginador '1 2 3 ... >>' y una fila fantasma de conteo), que si no se
    filtran entran como si fueran registros.
    """
    return exp.count("-") >= 3 and any(c.isdigit() for c in exp)


def parse_tabla(html_text, table_id, solapa):
    """Extrae los registros de una tabla (por id) y los etiqueta con la solapa."""
    soup = BeautifulSoup(html_text, "html.parser")
    tabla = soup.find("table", id=table_id)
    items = []
    if tabla is None:
        return items
    for tr in tabla.find_all("tr"):
        tds = tr.find_all("td")
        if len(tds) < 8:
            continue  # encabezado (th) o filas sin datos

        def celda(i):
            return tds[i].get_text(" ", strip=True) if i < len(tds) else ""

        item = {"solapa": solapa}
        for campo, idx in COL.items():
            item[campo] = celda(idx)

        if not _es_expediente(item["expediente"]):
            continue  # descarta paginador / pie de la grilla
        items.append(item)
    return items


def fetch_todo():
    """Descarga la pagina y devuelve (registros, notificaciones)."""
    r = requests.get(
        URL,
        timeout=45,
        headers={"User-Agent": "alertas-anmat-bot/1.0 (+github actions)"},
    )
    r.raise_for_status()
    html = r.content
    registros = parse_tabla(html, TABLA_REGISTROS, "Registros")
    notificaciones = parse_tabla(html, TABLA_NOTIFICACIONES, "Notificaciones")
    return registros, notificaciones


def clave(item):
    """Identificador unico e invisible para no repetir avisos.

    - Registros: se mantiene la clave = expediente (compatibilidad con el estado
      ya guardado en vistos_boletin.json; cambiarla reenviaria todo el historial).
    - Notificaciones: clave compuesta con prefijo NOTIF|, porque en esta solapa
      un mismo expediente puede repetirse para productos distintos (ej. 3083-10).
    """
    if item["solapa"] == "Notificaciones":
        return "NOTIF|" + "|".join([
            item["expediente"], item["nombre"], item["marca"], item["fecha"],
        ])
    if item["expediente"]:
        return item["expediente"]
    return "|".join([item["tramite"], item["nombre"], item["pm"], item["marca"]])


def dedup(items):
    """Colapsa filas repetidas (misma clave) conservando el orden."""
    vistos = set()
    out = []
    for it in items:
        k = clave(it)
        if k in vistos:
            continue
        vistos.add(k)
        out.append(it)
    return out


def load_seen():
    if STATE_FILE.exists():
        try:
            return set(json.loads(STATE_FILE.read_text(encoding="utf-8")))
        except Exception:
            return set()
    return None


def save_seen(claves):
    STATE_FILE.write_text(
        json.dumps(sorted(claves), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

# Tipo va primero para distinguir de un vistazo Registro vs Notificación.
# La columna "modelo" (Modelo/s) es la unica con ajuste de texto.
_HEADERS = [
    ("tipo", "Tipo", 14),
    ("tramite", "Trámite", 34),
    ("fecha", "Fecha", 12),
    ("razon", "Razón Social", 36),
    ("nombre", "Nombre", 34),
    ("marca", "Marca", 24),
    ("modelo", "Modelo/s", 60),
    ("pm", "PM", 14),
    ("expediente", "Expediente", 24),
]

# Etiqueta legible de la columna Tipo segun la solapa de origen.
_TIPO = {"Registros": "Registro", "Notificaciones": "Notificación"}

# La columna Modelo/s puede tener 30+ items y hacer filas gigantes. Limitamos
# la ALTURA visible a ~4 lineas; el texto completo queda en la celda (se ve
# expandiendo la fila) y el detalle fino esta siempre en el Boletin de ANMAT.
LINEAS_MODELO = 4
_ALTO_LINEA = 15          # alto aprox. de una linea (Calibri 11), en puntos
_ALTO_FILA = LINEAS_MODELO * _ALTO_LINEA


def _orden(it):
    """Ordena por Razón Social (A→Z, sin distinguir may/min), luego Tipo y Nombre."""
    return (
        (it.get("razon", "") or "").strip().upper(),
        it.get("tipo", ""),
        (it.get("nombre", "") or "").strip().upper(),
    )


def construir_excel(nuevos_reg, nuevos_notif, path):
    """Crea un .xlsx con UNA hoja que junta las dos solapas.

    Columna "Tipo" (Registro / Notificación) para diferenciarlas y ordenado
    por Razón Social, para ver juntas todas las filas de una misma empresa.
    """
    filas = []
    for it in nuevos_reg:
        filas.append({**it, "tipo": _TIPO["Registros"]})
    for it in nuevos_notif:
        filas.append({**it, "tipo": _TIPO["Notificaciones"]})
    filas.sort(key=_orden)

    wb = Workbook()
    ws = wb.active
    ws.title = "Novedades"

    header_fill = PatternFill("solid", fgColor="2E6DA4")
    header_font = Font(bold=True, color="FFFFFF")
    ws.append([h[1] for h in _HEADERS])
    for ci, (_k, _t, ancho) in enumerate(_HEADERS, 1):
        c = ws.cell(row=1, column=ci)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(ci)].width = ancho

    for it in filas:
        ws.append([it.get(k, "") for (k, _t, _w) in _HEADERS])

    # Ajuste de texto en Modelo/s (por nombre de encabezado, robusto al orden).
    col_modelo = [h[0] for h in _HEADERS].index("modelo") + 1
    for row in ws.iter_rows(min_row=2, min_col=col_modelo, max_col=col_modelo):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Altura fija de las filas de datos: muestra ~4 lineas de Modelo/s y oculta
    # el resto (la fila se puede expandir a mano para ver todo el texto).
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = _ALTO_FILA

    ws.freeze_panes = "A2"
    ultima_col = get_column_letter(len(_HEADERS))
    ws.auto_filter.ref = f"A1:{ultima_col}{max(ws.max_row, 1)}"
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Telegram
# ---------------------------------------------------------------------------

def send_telegram(text):
    url = f"https://api.telegram.org/bot{TOKEN}/sendMessage"
    resp = requests.post(
        url,
        data={"chat_id": CHAT_ID, "text": text, "parse_mode": "HTML"},
        timeout=30,
    )
    resp.raise_for_status()
    return resp.json()


def send_document(path, caption):
    url = f"https://api.telegram.org/bot{TOKEN}/sendDocument"
    with open(path, "rb") as f:
        resp = requests.post(
            url,
            data={"chat_id": CHAT_ID, "caption": caption, "parse_mode": "HTML"},
            files={"document": (os.path.basename(path), f)},
            timeout=120,
        )
    resp.raise_for_status()
    return resp.json()


def enviar_excel(nuevos_reg, nuevos_notif):
    """Arma el Excel con las novedades y lo manda como documento por Telegram."""
    hoy = datetime.now(TZ).strftime("%Y-%m-%d") if TZ else datetime.now().strftime("%Y-%m-%d")
    path = f"boletin_anmat_{hoy}.xlsx"
    construir_excel(nuevos_reg, nuevos_notif, path)
    total = len(nuevos_reg) + len(nuevos_notif)
    caption = (
        f"🆕 <b>Boletín ANMAT — {total} novedad(es)</b>\n"
        f"• Registros: {len(nuevos_reg)}\n"
        f"• Notificaciones: {len(nuevos_notif)}"
    )
    send_document(path, caption)


# ---------------------------------------------------------------------------
# Logica principal
# ---------------------------------------------------------------------------

def revisar():
    """Revisa las dos solapas, avisa lo nuevo (Excel) y devuelve (n_reg, n_notif)."""
    registros, notificaciones = fetch_todo()
    print(
        f"Boletin: {len(registros)} registros y {len(notificaciones)} "
        f"notificaciones en la pagina."
    )

    seen = load_seen()

    # Primera ejecucion absoluta (no existe el archivo de estado):
    # sembramos todo y avisamos que quedo activo, sin volcar el backlog.
    if seen is None:
        claves = {clave(r) for r in registros} | {clave(n) for n in notificaciones}
        save_seen(claves)
        send_telegram(
            "✅ <b>Aviso del Boletín ANMAT activado</b>\n\n"
            "Reviso las solapas <b>Registros</b> y <b>Notificaciones</b> y te "
            "aviso las novedades en un archivo Excel.\n"
            f"Ahora hay {len(registros)} registros y {len(notificaciones)} "
            "notificaciones recientes; de acá en más solo te aviso las nuevas."
        )
        print("Boletin: primera ejecucion, estado sembrado.")
        return 0, 0

    # ¿Ya venimos siguiendo Notificaciones? (la solapa se sumo despues).
    hay_estado_notif = any(k.startswith("NOTIF|") for k in seen)

    nuevos_reg = dedup([r for r in registros if clave(r) not in seen])

    if hay_estado_notif:
        nuevos_notif = dedup([n for n in notificaciones if clave(n) not in seen])
    else:
        # Primera vez con Notificaciones: sembrar sin avisar el backlog historico.
        nuevos_notif = []
        if notificaciones:
            send_telegram(
                "✅ <b>Notificaciones agregadas al aviso del Boletín ANMAT</b>\n\n"
                "Desde ahora también te aviso las nuevas declaraciones juradas "
                "PM I-II de la solapa <b>Notificaciones</b>, en el mismo Excel."
            )

    if nuevos_reg or nuevos_notif:
        enviar_excel(nuevos_reg, nuevos_notif)

    # Guardar estado con las claves actuales de las dos solapas (siembra
    # Notificaciones en esta corrida y mantiene fresca la ventana de vistos).
    seen |= {clave(r) for r in registros} | {clave(n) for n in notificaciones}
    save_seen(seen)

    print(f"Boletin: {len(nuevos_reg)} registros nuevos, {len(nuevos_notif)} notificaciones nuevas.")
    return len(nuevos_reg), len(nuevos_notif)


def main():
    if not TOKEN or not CHAT_ID:
        print("ERROR: faltan TELEGRAM_TOKEN o TELEGRAM_CHAT_ID.")
        sys.exit(1)
    revisar()


if __name__ == "__main__":
    main()
